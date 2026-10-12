"""Grounding rules, CV/JD loading, settings and usage logging."""

from __future__ import annotations

import docx
import pytest

from autoapply.config import ConfigError, load_settings
from autoapply.documents import InputError, load_cv, load_jd
from autoapply.grounding import (check_citations, check_company_names, check_numbers,
                                 check_unsupported_requirements, run_rule_checks)
from autoapply.models import Draft
from autoapply.usage import UsageLog, UsageRecord, format_table, summarise
from tests.conftest import SAMPLE_CV, SAMPLE_JD, clean_draft, fabricated_draft, parsed_jd


# --- grounding --------------------------------------------------------------

def test_clean_draft_passes_all_rules(cv) -> None:
    assert run_rule_checks(clean_draft(), cv, parsed_jd()) == []


def test_nonexistent_citation_flagged(cv) -> None:
    flags = check_citations(fabricated_draft(), cv)
    assert any("L99" in f.issue for f in flags)
    assert all(f.severity == "high" and f.source == "rule" for f in flags)


def test_uncited_claim_flagged(cv) -> None:
    d = Draft(cover_letter="x", claims=[{"sentence": "I am great.", "cv_lines": []}])
    assert check_citations(d, cv)[0].issue == "no CV line cited"


def test_invented_number_flagged(cv) -> None:
    flags = check_numbers(fabricated_draft(), cv)
    assert any("25" in f.issue for f in flags)


def test_number_present_in_cited_line_ok(cv) -> None:
    assert check_numbers(clean_draft(), cv) == []  # "40%" is in L07


def test_unsupported_jd_term_flagged(cv) -> None:
    flags = check_unsupported_requirements(fabricated_draft(), cv, parsed_jd())
    assert any(f.text == "SAP" for f in flags)


def test_acknowledged_gap_not_flagged(cv) -> None:
    d = Draft(cover_letter="I have not used SAP EWM yet.", gaps_acknowledged=["SAP EWM"])
    assert check_unsupported_requirements(d, cv, parsed_jd()) == []


def test_unknown_company_flagged_target_and_past_employer_ok(cv) -> None:
    d = Draft(cover_letter="Before Parcelio GmbH I advised Globex Ltd on routing. Nordwind Logistik GmbH appeals.")
    names = {f.text for f in check_company_names(d, cv, parsed_jd())}
    assert names == {"Globex Ltd"}


# --- documents ----------------------------------------------------------------

def test_cv_md_numbered_lines(cv) -> None:
    assert list(cv.lines)[0] == "L01"
    assert cv.lines["L01"] == "Alex Example"  # HTML comment header skipped
    assert "[L06]" in cv.numbered()
    assert "Power BI" in cv.lines["L06"]


def test_cv_docx_and_txt(tmp_path) -> None:
    d = docx.Document()
    d.add_paragraph("Jane Doe")
    d.add_paragraph("")
    d.add_paragraph("Built dashboards in Power BI", style="List Bullet")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "SQL"
    path = tmp_path / "cv.docx"
    d.save(path)
    cv = load_cv(path)
    assert list(cv.lines.values()) == ["Jane Doe", "Built dashboards in Power BI", "Skills | SQL"]
    txt = tmp_path / "cv.txt"
    txt.write_text("a\n\n  - b  \n", encoding="utf-8")
    assert list(load_cv(txt).lines.values()) == ["a", "b"]


@pytest.mark.parametrize("name,content,match", [
    ("cv.pdf", "x", "unsupported"), ("empty.md", "  \n", "empty"),
])
def test_cv_errors(tmp_path, name, content, match) -> None:
    p = tmp_path / name
    p.write_text(content, encoding="utf-8")
    with pytest.raises(InputError, match=match):
        load_cv(p)


def test_jd_refuses_urls() -> None:
    with pytest.raises(InputError, match="never fetches or scrapes"):
        load_jd("https://www.linkedin.com/jobs/view/123")


def test_jd_from_file_and_stdin(monkeypatch, tmp_path) -> None:
    assert "Supply Chain Data Analyst" in load_jd(str(SAMPLE_JD))
    import io
    monkeypatch.setattr("sys.stdin", io.StringIO("Pasted JD text"))
    assert load_jd("-") == "Pasted JD text"
    big = tmp_path / "big.md"
    big.write_text("x" * 30_000, encoding="utf-8")
    with pytest.raises(InputError, match="max"):
        load_jd(str(big))


# --- config ---------------------------------------------------------------------

def test_model_required(monkeypatch, tmp_path) -> None:
    monkeypatch.delenv("AUTOAPPLY_MODEL", raising=False)
    with pytest.raises(ConfigError, match="AUTOAPPLY_MODEL"):
        load_settings(tmp_path / "missing.env")


def test_settings_from_env(monkeypatch, tmp_path) -> None:
    monkeypatch.setenv("AUTOAPPLY_MODEL", "some-model")
    monkeypatch.setenv("AUTOAPPLY_EFFORT", "HIGH")
    monkeypatch.setenv("AUTOAPPLY_REFUSAL_FALLBACK", "false")
    monkeypatch.setenv("AUTOAPPLY_DATA_DIR", str(tmp_path))
    s = load_settings(None)
    assert (s.model, s.effort, s.refusal_fallback) == ("some-model", "high", False)
    assert s.checkpoint_db.parent == tmp_path
    monkeypatch.setenv("AUTOAPPLY_EFFORT", "turbo")
    with pytest.raises(ConfigError):
        load_settings(None)


# --- usage ------------------------------------------------------------------------

def test_usage_log_and_summary(tmp_path) -> None:
    log = UsageLog(tmp_path / "u.jsonl")
    for agent in ["parser", "writer", "writer"]:
        log.append(UsageRecord(run_id="r1", agent=agent, model="m", input_tokens=10, output_tokens=5))
    log.append(UsageRecord(run_id="r2", agent="parser", model="m", input_tokens=99, output_tokens=1))
    rows = summarise(log.records("r1"))
    assert [r["agent"] for r in rows] == ["parser", "writer", "TOTAL"]
    assert rows[1]["calls"] == 2 and rows[-1]["input_tokens"] == 30
    assert "| TOTAL | 3 | 30 |" in format_table(rows)
    assert len(log.records()) == 4


def test_sample_files_are_labelled_fictional() -> None:
    assert "FICTIONAL" in SAMPLE_CV.read_text(encoding="utf-8")
    assert "FICTIONAL" in SAMPLE_JD.read_text(encoding="utf-8")


def test_naming_the_role_is_not_an_unsupported_claim(cv) -> None:
    jd = parsed_jd().model_copy(update={"title": "Digital Procurement Manager",
                                        "must_haves": ["IT procurement background", "5+ years relevant experience"]})
    d = Draft(cover_letter="I would like to be considered for the Digital Procurement Manager role. "
                           "Relevant to your requirements: dashboards.")
    assert check_unsupported_requirements(d, cv, jd) == []
