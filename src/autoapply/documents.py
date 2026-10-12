"""Loading the CV and the job description from local files.

Design rules:
- The CV is split into numbered lines (``L01``, ``L02``...). Every claim the
  Writer makes must cite these IDs, which is what makes fabrication checkable.
- Job descriptions come from a local file or stdin only. URLs are refused on
  purpose: this tool does not scrape LinkedIn or job boards.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass
from pathlib import Path

import docx

MAX_JD_CHARS = 20_000
URL_PATTERN = re.compile(r"^\s*https?://", re.IGNORECASE)


class InputError(ValueError):
    """Raised for unusable CV or job-description input."""


@dataclass(frozen=True)
class CV:
    """A CV as numbered lines; ``lines`` maps line ID -> text."""

    lines: dict[str, str]
    source: str

    def numbered(self) -> str:
        """Render as ``[L01] text`` lines for prompts."""
        return "\n".join(f"[{k}] {v}" for k, v in self.lines.items())

    def text(self) -> str:
        """Plain text of the whole CV."""
        return "\n".join(self.lines.values())


def _read_docx(path: Path) -> list[str]:
    document = docx.Document(str(path))
    out = [p.text for p in document.paragraphs]
    for table in document.tables:  # some CV templates lay out sections in tables
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(dict.fromkeys(cells)))
    return out


def load_cv(path: str | Path) -> CV:
    """Load a CV from ``.md``, ``.txt`` or ``.docx`` into numbered lines."""
    path = Path(path)
    if not path.exists():
        raise InputError(f"CV file not found: {path}")
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        text = re.sub(r"<!--.*?-->", "", path.read_text(encoding="utf-8"), flags=re.DOTALL)  # drop HTML comments
        raw = text.splitlines()
    elif suffix == ".docx":
        raw = _read_docx(path)
    else:
        raise InputError(f"unsupported CV format {suffix!r}; use .md, .txt or .docx")
    cleaned = [re.sub(r"\s+", " ", line).strip(" \t-•*#") for line in raw]
    cleaned = [line for line in cleaned if line]
    if not cleaned:
        raise InputError(f"CV file is empty: {path}")
    width = max(2, len(str(len(cleaned))))
    return CV(lines={f"L{i:0{width}d}": line for i, line in enumerate(cleaned, start=1)}, source=str(path))


def load_jd(source: str) -> str:
    """Load a job description from a file path, or from stdin when ``source == '-'``.

    Raises:
        InputError: for URLs (no scraping), missing files, empty or oversized text.
    """
    if URL_PATTERN.match(source):
        raise InputError(
            "URLs are not accepted: autoapply never fetches or scrapes job boards. "
            "Copy the job description into a text file (or pipe it via --jd -)."
        )
    if source == "-":
        text = sys.stdin.read()
    else:
        path = Path(source)
        if not path.exists():
            raise InputError(f"job description file not found: {path}")
        text = path.read_text(encoding="utf-8")
    # Drop HTML comments: file headers (e.g. dataset attribution and labels) must never reach the models.
    text = re.sub(r"<!--.*?-->", "", text, flags=re.DOTALL).strip()
    if not text:
        raise InputError("job description is empty")
    if len(text) > MAX_JD_CHARS:
        raise InputError(f"job description is {len(text):,} characters; max is {MAX_JD_CHARS:,}")
    return text
